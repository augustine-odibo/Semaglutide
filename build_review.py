#!/usr/bin/env python3
"""
Build "Mechanisms of Action of Semaglutide and Tirzepatide" as Markdown and Word.

Usage:
    pip install matplotlib python-docx
    python build_review.py            # writes into ./review_output/

Outputs:
    review_output/fig1.png, fig2.png                 figures
    review_output/Semaglutide_Tirzepatide_Review.md   Markdown manuscript
    review_output/Semaglutide_Tirzepatide_Review.docx Word manuscript

Editing the text:
    All manuscript text lives in CONTENT at the bottom of this file.
    Cite with [[key]] or [[key1,key2]] using keys from CONTENT["REFS"];
    references are numbered automatically in order of first citation.
    Mark bold text with **double asterisks**.
"""
import os
import re

OUT = "review_output"

# ---------------------------------------------------------------- citations
CITE_RE = re.compile(r"\[\[([^\]]+)\]\]")


def walk_text(content):
    """Yield every piece of text in reading order, so citations number correctly."""
    for block in content["BODY"]:
        if "p" in block:
            yield block["p"]
        elif "bullets" in block:
            yield from block["bullets"]
        elif "table" in block:
            t = content["TABLES"][block["table"]]
            yield t["caption"]
            for row in t["rows"]:
                yield from row
            yield t["note"]
        elif "figure" in block:
            yield content["FIGURES"][block["figure"]]["caption"]


def number_refs(content):
    order = []
    for text in walk_text(content):
        for group in CITE_RE.findall(text):
            for key in (k.strip() for k in group.split(",")):
                if key not in content["REFS"]:
                    raise KeyError(f"Unknown reference key: {key}")
                if key not in order:
                    order.append(key)
    missing = [k for k in content["REFS"] if k not in order]
    if missing:
        raise ValueError(f"References never cited: {missing}")
    return {k: i + 1 for i, k in enumerate(order)}, order


def cite_label(group, numbers):
    ns = sorted({numbers[k.strip()] for k in group.split(",")})
    parts, start, prev = [], ns[0], ns[0]
    for n in ns[1:] + [None]:
        if n is not None and n == prev + 1:
            prev = n
            continue
        if start == prev:
            parts.append(str(start))
        elif prev == start + 1:
            parts.append(f"{start},{prev}")
        else:
            parts.append(f"{start}–{prev}")
        if n is not None:
            start = prev = n
    return "[" + ",".join(parts) + "]"


def resolve(text, numbers):
    return CITE_RE.sub(lambda m: cite_label(m.group(1), numbers), text)


def split_doi(ref):
    m = re.match(r"^(.*)doi:(\S+)$", ref)
    return (m.group(1), m.group(2)) if m else (ref, None)


def word_count(content):
    n = 0
    for block in content["BODY"]:
        texts = [block["p"]] if "p" in block else block.get("bullets", [])
        for t in texts:
            n += len(re.sub(r"\[\[[^\]]+\]\]|\*\*", "", t).split())
    return n


# ---------------------------------------------------------------- figures
def make_figures(out):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

    plt.rcParams["font.family"] = "DejaVu Sans"
    GLP, GIP, INK, MUTED = "#2B6CB0", "#C05621", "#1A202C", "#4A5568"

    # Figure 1: mechanism schematic
    fig, ax = plt.subplots(figsize=(10, 8), dpi=300)
    ax.set_xlim(0, 100)
    ax.set_ylim(-4, 76)
    ax.axis("off")

    def box(x, y, w, h, title, lines, edge, fc="white"):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.4,rounding_size=1.2",
                                    fc=fc, ec=edge, lw=1.6))
        ax.text(x + w / 2, y + h - 1.6, title, ha="center", va="top", fontsize=10.5,
                weight="bold", color=INK)
        for i, (t, c) in enumerate(lines):
            ax.text(x + 1.2, y + h - 5.6 - i * 2.9, t, ha="left", va="top", fontsize=8.3, color=c)

    box(35, 43, 30, 12, "Semaglutide",
        [("Selective GLP-1R agonist", GLP), ("C18 diacid, albumin-bound; t½ ≈ 1 wk", MUTED)],
        GLP, fc="#EBF4FF")
    box(35, 24, 30, 15, "Tirzepatide",
        [("Dual GIPR / GLP-1R agonist", INK), ("GIPR-favoured (imbalanced)", GIP),
         ("cAMP-biased at GLP-1R", GLP)], GIP, fc="#FFF5EB")
    organs = [
        (1, 57, 31, 17, "Brain", [("↓ food intake, ↑ satiety", GLP), ("↓ cravings / food reward", GLP),
                                  ("+ GIPR: added anorexia,", GIP), ("   ↓ nausea and emesis", GIP)]),
        (68, 57, 31, 17, "Pancreatic islet", [("↑ glucose-dependent insulin", GLP),
                                              ("↓ glucagon in hyperglycaemia", GLP),
                                              ("+ GIPR: additional insulin;", GIP),
                                              ("   glucagon at low glucose", GIP)]),
        (1, 31, 31, 15, "Stomach and gut", [("↓ gastric emptying", GLP),
                                            ("   (wanes with chronic use)", MUTED),
                                            ("GI adverse effects", MUTED)]),
        (68, 31, 31, 15, "Adipose tissue", [("+ GIPR: ↑ blood flow,", GIP), ("   lipid buffering", GIP),
                                            ("+ GIPR: ↑ insulin sensitivity", GIP)]),
        (1, 3, 31, 15, "Heart and vessels", [("↓ MACE (SELECT, SUSTAIN-6)", GLP),
                                             ("↓ HF events in HFpEF (SUMMIT)", INK),
                                             ("↓ BP, ↓ CRP; ↑ heart rate", GLP)]),
        (35, 3, 30, 15, "Liver", [("Indirect: ↓ steatosis,", MUTED), ("   MASH resolution", MUTED),
                                  ("(little hepatocyte GLP-1R)", MUTED)]),
        (68, 3, 31, 15, "Kidney", [("↓ kidney disease events (FLOW)", GLP), ("↓ albuminuria", GLP),
                                   ("slower eGFR decline", GLP)]),
    ]
    for o in organs:
        box(*o[:5], o[5], "#A0AEC0")
    for sx, sy, tx, ty in [(35, 50, 32.4, 64), (65, 50, 67.6, 64), (35, 30, 32.4, 38), (65, 30, 67.6, 38),
                           (38, 23.6, 32.4, 14), (50, 23.6, 50, 18.4), (62, 23.6, 67.6, 14)]:
        ax.add_patch(FancyArrowPatch((sx, sy), (tx, ty), arrowstyle="-|>", mutation_scale=12,
                                     lw=1.3, color="#718096"))
    for i, (c, label) in enumerate([(GLP, "GLP-1R-mediated (both drugs)"),
                                    (GIP, "GIPR-mediated (tirzepatide only)"),
                                    (MUTED, "Indirect or secondary")]):
        ax.plot([12 + i * 28], [-2.2], marker="s", color=c, ms=7)
        ax.text(13.5 + i * 28, -2.2, label, fontsize=8.2, va="center", color=INK)
    fig.savefig(os.path.join(out, "fig1.png"), bbox_inches="tight", facecolor="white")
    plt.close(fig)

    # Figure 2: weight loss across trials (values from the cited trial reports)
    rows = [
        ("Retatrutide 12 mg (phase 2, 48 wk)", "Placebo", -24.2, -2.1),
        ("CagriSema 2.4/2.4 mg (REDEFINE 1, 68 wk)", "Placebo", -20.4, -3.0),
        ("Orforglipron 36 mg oral (ATTAIN-1, 72 wk)", "Placebo", -11.2, -2.1),
        ("Tirzepatide MTD (SURMOUNT-5, 72 wk)", "Semaglutide MTD", -20.2, -13.7),
        ("Tirzepatide 15 mg (SURMOUNT-1, 72 wk)", "Placebo", -20.9, -3.1),
        ("Semaglutide 2.4 mg (STEP 1, 68 wk)", "Placebo", -14.9, -2.4),
    ]
    fig, ax = plt.subplots(figsize=(9, 4.6), dpi=300)
    for i, (label, comp, active, control) in enumerate(rows):
        col = GIP if "Tirzepatide" in label else (GLP if "Semaglutide" in label else MUTED)
        ax.plot([control, active], [i, i], color="#CBD5E0", lw=3, zorder=1)
        ax.scatter(control, i, s=60, color="#A0AEC0" if comp == "Placebo" else GLP, zorder=2,
                   edgecolor="white")
        ax.scatter(active, i, s=80, color=col, zorder=3, edgecolor="white")
        ax.text(active - 0.6, i + 0.28, f"{active:.1f}%", ha="center", fontsize=8, color=INK)
        ax.text(control + 0.6, i + 0.28, f"{control:.1f}%" + ("" if comp == "Placebo" else " (sema)"),
                ha="center", fontsize=8, color=MUTED)
    ax.set_yticks(range(len(rows)))
    ax.set_yticklabels([r[0] for r in rows], fontsize=8.5)
    ax.axhline(2.5, color="#E2E8F0", lw=1, ls="--")
    ax.text(-27.5, 2.62, "Approved agents", fontsize=7.5, color=MUTED)
    ax.text(-27.5, 2.38, "Investigational agents", fontsize=7.5, color=MUTED, va="top")
    ax.set_xlim(-28, 1)
    ax.set_ylim(-0.6, 5.8)
    ax.set_xlabel("Mean change in body weight from baseline (%)", fontsize=9)
    for s in ["top", "right", "left"]:
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.grid(axis="x", color="#EDF2F7")
    for c, label in [("#A0AEC0", "Placebo"), (GLP, "Semaglutide"), (GIP, "Tirzepatide"), (MUTED, "Other agent")]:
        ax.scatter([], [], color=c, label=label)
    ax.legend(frameon=False, fontsize=7.5, loc="lower left", ncol=4, bbox_to_anchor=(0, -0.28))
    fig.savefig(os.path.join(out, "fig2.png"), bbox_inches="tight", facecolor="white")
    plt.close(fig)


# ---------------------------------------------------------------- Markdown
def md_escape_cell(s):
    return s.replace("|", "\\|")


def build_markdown(content, numbers, order, path):
    C, R = content, (lambda t: resolve(t, numbers))
    n_words = round(word_count(C) / 50) * 50
    L = [f"# {C['TITLE']}", "",
         "*Review article*", "",
         "**Authors:** [Author name(s)] · [Affiliation(s)]  ",
         "**Corresponding author:** [Name, address, email]  ",
         f"**Running title:** {C['RUNNING']}  ",
         f"**Main text:** approximately {n_words} words · 3 tables · 2 figures · {len(order)} references", "",
         "---", "", "## Abstract", ""]
    L += [R(a) + "\n" for a in C["ABSTRACT"]]
    L += [f"**Keywords:** {C['KEYWORDS']}", "", "> **Key points**", ">"]
    L += [f"> - {R(k)}" for k in C["KEYPOINTS"]]
    L += ["", "### Abbreviations", "",
          "; ".join(f"**{a}**, {d}" for a, d in C["ABBREV"]) + ".", ""]
    for b in C["BODY"]:
        if "h1" in b:
            L += [f"## {b['h1']}", ""]
        elif "h2" in b:
            L += [f"### {b['h2']}", ""]
        elif "p" in b:
            L += [R(b["p"]), ""]
        elif "bullets" in b:
            L += [f"- {R(x)}" for x in b["bullets"]] + [""]
        elif "table" in b:
            t = C["TABLES"][b["table"]]
            L += [f"**{R(t['caption'])}**", "",
                  "| " + " | ".join(t["head"]) + " |",
                  "|" + "|".join("---" for _ in t["head"]) + "|"]
            L += ["| " + " | ".join(md_escape_cell(R(c)) for c in row) + " |" for row in t["rows"]]
            L += ["", f"*{R(t['note'])}*", ""]
        elif "figure" in b:
            f = C["FIGURES"][b["figure"]]
            L += [f"![{f['file']}]({f['file']})", "", R(f["caption"]), ""]
    L += ["## Declarations", "",
          "**Funding:** [State funding sources, or “None.”]  ",
          "**Competing interests:** [Declare any relationships with manufacturers of incretin-based therapies, or “None.”]  ",
          "**Author contributions:** [Describe contributions using the CRediT taxonomy.]  ",
          "**Use of AI tools:** [Disclose any use of AI-assisted drafting according to the target journal’s policy.]", "",
          "## References", ""]
    for i, k in enumerate(order, 1):
        text, doi = split_doi(C["REFS"][k])
        L.append(f"{i}. {text}" + (f"doi:[{doi}](https://doi.org/{doi})" if doi else ""))
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L) + "\n")


# ---------------------------------------------------------------- Word
def build_docx(content, numbers, order, out, path):
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches

    C = content
    ACCENT = RGBColor(0x1F, 0x4E, 0x79)
    doc = Document()

    # page + base styles
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, italic in (("Heading 1", 14, False), ("Heading 2", 12, True)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold, st.font.italic = "Times New Roman", Pt(size), True, italic
        st.font.color.rgb = ACCENT if name == "Heading 1" else RGBColor(0x26, 0x26, 0x26)
        rfonts = st.element.rPr.find(qn("w:rFonts"))
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), "Times New Roman")
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
        st.paragraph_format.space_before = Pt(16 if name == "Heading 1" else 12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True

    # header / footer with page number
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run(C["RUNNING"])
    r.italic, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(0x7F, 0x7F, 0x7F)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        fr._r.append(el)

    def add_runs(par, text, size=None, bold=False, italic=False, color=None):
        text = resolve(text, numbers)
        for part in re.split(r"(\*\*[^*]+\*\*)", text):
            if not part:
                continue
            is_bold = part.startswith("**")
            run = par.add_run(part[2:-2] if is_bold else part)
            run.bold = bold or is_bold
            run.italic = italic
            if size:
                run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = color
        return par

    def para(text, justify=True, **kw):
        p = doc.add_paragraph()
        if justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return add_runs(p, text, **kw)

    def shade(cell, hex_fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    def page_break():
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def hyperlink(par, url, text, size):
        part = par.part
        r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                              is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), r_id)
        new_run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        for tag, val in (("w:color", "0563C1"), ("w:u", "single"), ("w:sz", str(int(size * 2)))):
            e = OxmlElement(tag)
            e.set(qn("w:val"), val)
            rpr.append(e)
        new_run.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        new_run.append(t)
        h.append(new_run)
        par._p.append(h)

    # ---- title page
    p = doc.add_paragraph()
    r = p.add_run("REVIEW ARTICLE")
    r.font.size, r.font.color.rgb = Pt(10), RGBColor(0x59, 0x59, 0x59)
    p = doc.add_paragraph()
    r = p.add_run(C["TITLE"])
    r.bold, r.font.size, r.font.color.rgb = True, Pt(18), ACCENT
    n_words = round(word_count(C) / 50) * 50
    para("[Author name(s)] · [Affiliation(s)]", justify=False)
    para("**Corresponding author:** [Name, address, email]", justify=False)
    para(f"**Running title:** {C['RUNNING']}", justify=False)
    para(f"**Main text:** approximately {n_words} words · 3 tables · 2 figures · {len(order)} references", justify=False)
    doc.add_heading("Abstract", level=1)
    for a in C["ABSTRACT"]:
        para(a)
    para(f"**Keywords:** {C['KEYWORDS']}", justify=False)
    page_break()

    # ---- key points box + abbreviations
    box = doc.add_table(rows=1, cols=1)
    cell = box.rows[0].cells[0]
    shade(cell, "EAF1F8")
    add_runs(cell.paragraphs[0], "Key points", size=11, bold=True, color=ACCENT)
    for k in C["KEYPOINTS"]:
        add_runs(cell.add_paragraph(style="List Bullet"), k, size=10)
    doc.add_paragraph()
    doc.add_heading("Abbreviations", level=2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for i, (a, d) in enumerate(C["ABBREV"]):
        r = p.add_run(a)
        r.bold, r.font.size = True, Pt(10)
        p.add_run(f", {d}{'; ' if i < len(C['ABBREV']) - 1 else '.'}").font.size = Pt(10)
    page_break()

    # ---- body
    for b in C["BODY"]:
        if "h1" in b:
            doc.add_heading(b["h1"], level=1)
        elif "h2" in b:
            doc.add_heading(b["h2"], level=2)
        elif "p" in b:
            para(b["p"])
        elif "bullets" in b:
            for x in b["bullets"]:
                add_runs(doc.add_paragraph(style="List Bullet"), x)
        elif "table" in b:
            t = C["TABLES"][b["table"]]
            cap = add_runs(doc.add_paragraph(), t["caption"], size=10, bold=True)
            cap.paragraph_format.keep_with_next = True
            tbl = doc.add_table(rows=1, cols=len(t["head"]))
            tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            widths = [Inches(w / 1440) for w in t["widths"]]
            for i, h in enumerate(t["head"]):
                c = tbl.rows[0].cells[i]
                shade(c, "1F4E79")
                add_runs(c.paragraphs[0], h, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            for ri, row in enumerate(t["rows"]):
                cells = tbl.add_row().cells
                for i, v in enumerate(row):
                    if ri % 2:
                        shade(cells[i], "F2F6FA")
                    add_runs(cells[i].paragraphs[0], v, size=9)
            for row in tbl.rows:
                for i, c in enumerate(row.cells):
                    c.width = widths[i]
                    c.paragraphs[0].paragraph_format.space_after = Pt(0)
                    c.paragraphs[0].paragraph_format.line_spacing = 1.0
            add_runs(doc.add_paragraph(), t["note"], size=8.5, italic=True)
        elif "figure" in b:
            f = C["FIGURES"][b["figure"]]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(os.path.join(out, f["file"]), width=Inches(6.4))
            para(f["caption"], size=9)

    # ---- declarations + references
    doc.add_heading("Declarations", level=1)
    for line in ("**Funding:** [State funding sources, or “None.”]",
                 "**Competing interests:** [Declare any relationships with manufacturers of incretin-based therapies, or “None.”]",
                 "**Author contributions:** [Describe contributions using the CRediT taxonomy.]",
                 "**Use of AI tools:** [Disclose any use of AI-assisted drafting according to the target journal’s policy.]"):
        para(line, justify=False)
    doc.add_heading("References", level=1)
    for i, k in enumerate(order, 1):
        text, doi = split_doi(C["REFS"][k])
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(f"{i}.\t{text}doi:" if doi else f"{i}.\t{text}")
        r.font.size = Pt(9.5)
        if doi:
            hyperlink(p, f"https://doi.org/{doi}", doi, 9.5)

    # python-docx's default template omits a required zoom attribute; add it so the file validates
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    doc.core_properties.title = C["TITLE"]
    doc.save(path)


# ---------------------------------------------------------------- main
def main():
    os.makedirs(OUT, exist_ok=True)
    numbers, order = number_refs(CONTENT)
    make_figures(OUT)
    md_path = os.path.join(OUT, "Semaglutide_Tirzepatide_Review.md")
    docx_path = os.path.join(OUT, "Semaglutide_Tirzepatide_Review.docx")
    build_markdown(CONTENT, numbers, order, md_path)
    build_docx(CONTENT, numbers, order, OUT, docx_path)
    print(f"Wrote {md_path}\nWrote {docx_path}\n{len(order)} references, ~{word_count(CONTENT)} words")


# ---------------------------------------------------------------- content
CONTENT = {'REFS': {'nauck86': 'Nauck M, Stöckmann F, Ebert R, Creutzfeldt W. Reduced incretin effect in type 2 '
                     '(non-insulin-dependent) diabetes. Diabetologia. 1986;29(1):46-52. '
                     'doi:10.1007/BF02427280',
          'nauck93': 'Nauck MA, Heimesaat MM, Orskov C, Holst JJ, Ebert R, Creutzfeldt W. Preserved incretin '
                     'activity of glucagon-like peptide 1 [7-36 amide] but not of synthetic human gastric '
                     'inhibitory polypeptide in patients with type-2 diabetes mellitus. J Clin Invest. '
                     '1993;91(1):301-7. doi:10.1172/JCI116186',
          'drucker18': 'Drucker DJ. Mechanisms of action and therapeutic application of glucagon-like '
                       'peptide-1. Cell Metab. 2018;27(4):740-56. doi:10.1016/j.cmet.2018.03.001',
          'lau15': 'Lau J, Bloch P, Schäffer L, et al. Discovery of the once-weekly glucagon-like peptide-1 '
                   '(GLP-1) analogue semaglutide. J Med Chem. 2015;58(18):7370-80. '
                   'doi:10.1021/acs.jmedchem.5b00726',
          'marso16': 'Marso SP, Bain SC, Consoli A, et al. Semaglutide and cardiovascular outcomes in '
                     'patients with type 2 diabetes. N Engl J Med. 2016;375(19):1834-44. '
                     'doi:10.1056/NEJMoa1607141',
          'coskun18': 'Coskun T, Sloop KW, Loghin C, et al. LY3298176, a novel dual GIP and GLP-1 receptor '
                      'agonist for the treatment of type 2 diabetes mellitus: from discovery to clinical '
                      'proof of concept. Mol Metab. 2018;18:3-14. doi:10.1016/j.molmet.2018.09.009',
          'willard20': 'Willard FS, Douros JD, Gabe MBN, et al. Tirzepatide is an imbalanced and biased dual '
                       'GIP and GLP-1 receptor agonist. JCI Insight. 2020;5(17):e140532. '
                       'doi:10.1172/jci.insight.140532',
          'sun22': 'Sun B, Willard FS, Feng D, et al. Structural determinants of dual incretin receptor '
                   'agonism by tirzepatide. Proc Natl Acad Sci U S A. 2022;119(13):e2116506119. '
                   'doi:10.1073/pnas.2116506119',
          'heise22': 'Heise T, Mari A, DeVries JH, et al. Effects of subcutaneous tirzepatide versus placebo '
                     'or semaglutide on pancreatic islet function and insulin sensitivity in adults with '
                     'type 2 diabetes: a multicentre, randomised, double-blind, parallel-arm, phase 1 '
                     'clinical trial. Lancet Diabetes Endocrinol. 2022;10(6):418-29. '
                     'doi:10.1016/S2213-8587(22)00085-7',
          'samms20': 'Samms RJ, Coghlan MP, Sloop KW. How may GIP enhance the therapeutic efficacy of GLP-1? '
                     'Trends Endocrinol Metab. 2020;31(6):410-21. doi:10.1016/j.tem.2020.02.006',
          'samms21': 'Samms RJ, Christe ME, Collins KA, et al. GIPR agonism mediates weight-independent '
                     'insulin sensitization by tirzepatide in obese mice. J Clin Invest. '
                     '2021;131(12):e146353. doi:10.1172/JCI146353',
          'friedrichsen21': 'Friedrichsen M, Breitschaft A, Tadayon S, Wizert A, Skovgaard D. The effect of '
                            'semaglutide 2.4 mg once weekly on energy intake, appetite, control of eating, '
                            'and gastric emptying in adults with obesity. Diabetes Obes Metab. '
                            '2021;23(3):754-62. doi:10.1111/dom.14280',
          'gabery20': 'Gabery S, Salinas CG, Paulsen SJ, et al. Semaglutide lowers body weight in rodents '
                      'via distributed neural pathways. JCI Insight. 2020;5(6):e133429. '
                      'doi:10.1172/jci.insight.133429',
          'secher14': 'Secher A, Jelsing J, Baquero AF, et al. The arcuate nucleus mediates GLP-1 receptor '
                      'agonist liraglutide-dependent weight loss. J Clin Invest. 2014;124(10):4473-88. '
                      'doi:10.1172/JCI75276',
          'hendershot25': 'Hendershot CS, Bremmer MP, Paladino MB, et al. Once-weekly semaglutide in adults '
                          'with alcohol use disorder: a randomized clinical trial. JAMA Psychiatry. '
                          '2025;82(4):395-405. doi:10.1001/jamapsychiatry.2024.4789',
          'zhang21': 'Zhang Q, Delessa CT, Augustin R, et al. The glucose-dependent insulinotropic '
                     'polypeptide (GIP) regulates body weight and food intake via CNS-GIPR signaling. Cell '
                     'Metab. 2021;33(4):833-44.e5. doi:10.1016/j.cmet.2021.01.015',
          'adriaenssens23': 'Adriaenssens A, Broichhagen J, de Bray A, et al. Hypothalamic and brainstem '
                            'glucose-dependent insulinotropic polypeptide receptor neurons employ distinct '
                            'mechanisms to affect feeding. JCI Insight. 2023;8(10):e164921. '
                            'doi:10.1172/jci.insight.164921',
          'borner21': 'Borner T, Geisler CE, Fortin SM, et al. GIP receptor agonism attenuates GLP-1 '
                      'receptor agonist-induced nausea and emesis in preclinical models. Diabetes. '
                      '2021;70(11):2545-53. doi:10.2337/db21-0459',
          'sanyal25': 'Sanyal AJ, Newsome PN, Kliers I, et al. Phase 3 trial of semaglutide in metabolic '
                      'dysfunction-associated steatohepatitis. N Engl J Med. 2025;392(21):2089-99. '
                      'doi:10.1056/NEJMoa2413258',
          'loomba24': 'Loomba R, Hartman ML, Lawitz EJ, et al. Tirzepatide for metabolic '
                      'dysfunction-associated steatohepatitis with liver fibrosis. N Engl J Med. '
                      '2024;391(4):299-310. doi:10.1056/NEJMoa2401943',
          'lincoff23': 'Lincoff AM, Brown-Frandsen K, Colhoun HM, et al. Semaglutide and cardiovascular '
                       'outcomes in obesity without diabetes. N Engl J Med. 2023;389(24):2221-32. '
                       'doi:10.1056/NEJMoa2307563',
          'deanfield25': 'Deanfield J, Lincoff AM, Kahn SE, et al. Semaglutide and cardiovascular outcomes '
                         'by baseline and changes in adiposity measurements: a prespecified analysis of the '
                         'SELECT trial. Lancet. 2025;406(10516):2257-68. doi:10.1016/S0140-6736(25)01375-3',
          'kosiborod23': 'Kosiborod MN, Abildstrøm SZ, Borlaug BA, et al. Semaglutide in patients with heart '
                         'failure with preserved ejection fraction and obesity. N Engl J Med. '
                         '2023;389(12):1069-84. doi:10.1056/NEJMoa2306963',
          'packer25': 'Packer M, Zile MR, Kramer CM, et al. Tirzepatide for heart failure with preserved '
                      'ejection fraction and obesity. N Engl J Med. 2025;392(5):427-37. '
                      'doi:10.1056/NEJMoa2410027',
          'nicholls25': 'Nicholls SJ, Pavo I, Bhatt DL, et al. Cardiovascular outcomes with tirzepatide '
                        'versus dulaglutide in type 2 diabetes. N Engl J Med. 2025;393(24):2409-20. '
                        'doi:10.1056/NEJMoa2505928',
          'perkovic24': 'Perkovic V, Tuttle KR, Rossing P, et al. Effects of semaglutide on chronic kidney '
                        'disease in patients with type 2 diabetes. N Engl J Med. 2024;391(2):109-21. '
                        'doi:10.1056/NEJMoa2403347',
          'killion18': 'Killion EA, Wang J, Yie J, et al. Anti-obesity effects of GIPR antagonists alone and '
                       'in combination with GLP-1R agonists in preclinical models. Sci Transl Med. '
                       '2018;10(472):eaat3392. doi:10.1126/scitranslmed.aat3392',
          'killion20': 'Killion EA, Chen M, Falsey JR, et al. Chronic glucose-dependent insulinotropic '
                       'polypeptide receptor (GIPR) agonism desensitizes adipocyte GIPR activity mimicking '
                       'functional GIPR antagonism. Nat Commun. 2020;11(1):4981. '
                       'doi:10.1038/s41467-020-18751-8',
          'veniant24': 'Véniant MM, Lu SC, Atangan L, et al. A GIPR antagonist conjugated to GLP-1 analogues '
                       'promotes weight loss with improved metabolic parameters in preclinical and phase 1 '
                       'settings. Nat Metab. 2024;6(2):290-303. doi:10.1038/s42255-023-00966-w',
          'frias21': 'Frías JP, Davies MJ, Rosenstock J, et al. Tirzepatide versus semaglutide once weekly '
                     'in patients with type 2 diabetes. N Engl J Med. 2021;385(6):503-15. '
                     'doi:10.1056/NEJMoa2107519',
          'wilding21': 'Wilding JPH, Batterham RL, Calanna S, et al. Once-weekly semaglutide in adults with '
                       'overweight or obesity. N Engl J Med. 2021;384(11):989-1002. '
                       'doi:10.1056/NEJMoa2032183',
          'jastreboff22': 'Jastreboff AM, Aronne LJ, Ahmad NN, et al. Tirzepatide once weekly for the '
                          'treatment of obesity. N Engl J Med. 2022;387(3):205-16. doi:10.1056/NEJMoa2206038',
          'aronne25': 'Aronne LJ, Horn DB, le Roux CW, et al. Tirzepatide as compared with semaglutide for '
                      'the treatment of obesity. N Engl J Med. 2025;393(1):26-36. doi:10.1056/NEJMoa2416394',
          'wilding22': 'Wilding JPH, Batterham RL, Davies M, et al. Weight regain and cardiometabolic '
                       'effects after withdrawal of semaglutide: the STEP 1 trial extension. Diabetes Obes '
                       'Metab. 2022;24(8):1553-64. doi:10.1111/dom.14725',
          'aronne24': 'Aronne LJ, Sattar N, Horn DB, et al. Continued treatment with tirzepatide for '
                      'maintenance of weight reduction in adults with obesity: the SURMOUNT-4 randomized '
                      'clinical trial. JAMA. 2024;331(1):38-48. doi:10.1001/jama.2023.24945',
          'karakasis24': 'Karakasis P, Patoulias D, Fragakis N, Mantzoros CS. Effect of glucagon-like '
                         'peptide-1 receptor agonists and co-agonists on body composition: systematic review '
                         'and network meta-analysis. Metabolism. 2025;164:156113. '
                         'doi:10.1016/j.metabol.2024.156113',
          'stefanakis24': 'Stefanakis K, Kokkorakis M, Mantzoros CS. The impact of weight loss on fat-free '
                          'mass, muscle, bone and hematopoiesis health: implications for emerging '
                          'pharmacotherapies aiming at fat reduction and lean mass preservation. Metabolism. '
                          '2024;161:156057. doi:10.1016/j.metabol.2024.156057',
          'bjerre10': 'Bjerre Knudsen L, Madsen LW, Andersen S, et al. Glucagon-like peptide-1 receptor '
                      'agonists activate rodent thyroid C-cells causing calcitonin release and C-cell '
                      'proliferation. Endocrinology. 2010;151(4):1473-86. doi:10.1210/en.2009-1272',
          'hathaway24': 'Hathaway JT, Shah MP, Hathaway DB, et al. Risk of nonarteritic anterior ischemic '
                        'optic neuropathy in patients prescribed semaglutide. JAMA Ophthalmol. '
                        '2024;142(8):732-9. doi:10.1001/jamaophthalmol.2024.2296',
          'jastreboff23': 'Jastreboff AM, Kaplan LM, Frías JP, et al. Triple-hormone-receptor agonist '
                          'retatrutide for obesity - a phase 2 trial. N Engl J Med. 2023;389(6):514-26. '
                          'doi:10.1056/NEJMoa2301972',
          'garvey25': 'Garvey WT, Blüher M, Osorto Contreras CK, et al. Coadministered cagrilintide and '
                      'semaglutide in adults with overweight or obesity. N Engl J Med. 2025;393(7):635-47. '
                      'doi:10.1056/NEJMoa2502081',
          'wharton25': 'Wharton S, Aronne LJ, Stefanski A, et al. Orforglipron, an oral small-molecule GLP-1 '
                       'receptor agonist for obesity treatment. N Engl J Med. 2025;393(18):1796-806. '
                       'doi:10.1056/NEJMoa2511774'},
 'TITLE': 'Mechanisms of Action of Semaglutide and Tirzepatide: From Receptor Pharmacology to Clinical '
          'Outcomes',
 'RUNNING': 'Mechanisms of semaglutide and tirzepatide',
 'ABSTRACT': ['Semaglutide (Ozempic, Wegovy) and tirzepatide (Mounjaro, Zepbound) are once-weekly '
              'incretin-based peptides that produce glucose lowering and weight loss approaching that '
              'achieved with bariatric surgery. Although both mimic gut-derived incretin hormones, they '
              'differ fundamentally in receptor pharmacology. Semaglutide is a selective, albumin-bound '
              'glucagon-like peptide-1 receptor (GLP-1R) agonist with a half-life of about one week. '
              'Tirzepatide is a single GIP-based peptide that engages the glucose-dependent insulinotropic '
              'polypeptide receptor (GIPR) more than GLP-1R and acts at GLP-1R as a cAMP-biased agonist with '
              'reduced receptor internalisation. Both drugs restore glucose-dependent insulin secretion and '
              'suppress glucagon, but most of their effect on body weight arises from reduced energy intake '
              'driven by GLP-1R signalling at circumventricular, hypothalamic and hindbrain sites; slowing '
              'of gastric emptying is largely lost with chronic dosing. The GIPR component of tirzepatide '
              'adds weight-independent insulin sensitisation through adipose tissue, central anorexigenic '
              'signalling and, in preclinical models, suppression of GLP-1R agonist-induced emesis. These '
              'differences are consistent with the superiority of tirzepatide over semaglutide for weight '
              'loss (20.2% vs 13.7% at 72 weeks in SURMOUNT-5). Outcome trials show that semaglutide reduces '
              'major adverse cardiovascular events and kidney disease progression, largely independently of '
              'weight loss, while tirzepatide reduces heart failure events in obesity-related heart failure '
              'with preserved ejection fraction. We synthesise the molecular, physiological and clinical '
              'evidence for these mechanisms, examine the unresolved paradox that both GIPR agonism and '
              'antagonism enhance weight loss, discuss mechanism-linked adverse effects, and outline the '
              'multi-receptor and oral agents now emerging.'],
 'KEYWORDS': 'semaglutide; tirzepatide; GLP-1 receptor; GIP receptor; incretin; obesity; type 2 diabetes; '
             'biased agonism',
 'KEYPOINTS': ['Semaglutide is a selective GLP-1R agonist; tirzepatide is an imbalanced (GIPR-favoured), '
               'cAMP-biased dual GIPR/GLP-1R agonist.',
               'Weight loss with both drugs is driven mainly by reduced food intake via GLP-1R signalling in '
               'circumventricular, hypothalamic and hindbrain circuits, not by sustained gastric slowing.',
               'GIPR agonism contributes weight-independent insulin sensitisation through adipose tissue and '
               'may improve tolerability by suppressing hindbrain emetic signalling.',
               "Semaglutide's cardiovascular benefit in SELECT was largely independent of the amount of "
               'weight lost; about one-third was statistically mediated by waist reduction.',
               'Both drugs require continued treatment: about two-thirds of lost weight returns within a '
               'year of stopping.',
               'Both GIPR agonism and GIPR antagonism enhance GLP-1R-driven weight loss; receptor '
               'desensitisation is one leading explanation of this paradox.'],
 'ABBREV': [['AgRP', 'agouti-related peptide'],
            ['AP', 'area postrema'],
            ['CVO', 'circumventricular organ'],
            ['DPP-4', 'dipeptidyl peptidase-4'],
            ['eGFR', 'estimated glomerular filtration rate'],
            ['GIP(R)', 'glucose-dependent insulinotropic polypeptide (receptor)'],
            ['GLP-1(R)', 'glucagon-like peptide-1 (receptor)'],
            ['HFpEF', 'heart failure with preserved ejection fraction'],
            ['MACE', 'major adverse cardiovascular events'],
            ['MASH', 'metabolic dysfunction-associated steatohepatitis'],
            ['NAION', 'non-arteritic anterior ischaemic optic neuropathy'],
            ['NTS', 'nucleus of the solitary tract'],
            ['POMC', 'pro-opiomelanocortin'],
            ['T2D', 'type 2 diabetes']],
 'BODY': [{'h1': '1. Introduction'},
          {'p': 'Semaglutide and tirzepatide have changed the pharmacological treatment of type 2 diabetes '
                '(T2D) and obesity. Both are once-weekly subcutaneous peptides derived from the incretin '
                'hormones, yet they engage different receptors and signal differently once bound. '
                'Understanding these differences explains why tirzepatide produces greater glucose lowering '
                'and weight loss, why both drugs share most of their adverse effects, and why their benefits '
                'wane when treatment stops.'},
          {'p': 'The therapeutic rationale rests on the incretin effect: oral glucose elicits far more '
                'insulin secretion than an isoglycaemic intravenous glucose infusion. In healthy people, '
                'incretin factors account for roughly 73% of the insulin response to oral glucose, but only '
                'about 36% in people with T2D [[nauck86]]. Native glucagon-like peptide-1 (GLP-1) and '
                'glucose-dependent insulinotropic polypeptide (GIP) are cleared within minutes '
                '[[drucker18]], so drug development focused on stable, long-acting receptor agonists. '
                'Exenatide and liraglutide established the class; semaglutide extended dosing to once '
                'weekly; and tirzepatide added a second receptor.'},
          {'p': 'This review integrates molecular, physiological and clinical evidence on how each drug '
                'works. We begin with incretin biology and receptor pharmacology, then examine actions on '
                'the islet, brain, gut, adipose tissue, liver, heart and kidney, before linking mechanism to '
                'efficacy, durability and safety. We close with the unresolved GIP agonist–antagonist '
                'paradox and the next generation of incretin-based therapies. All cited references were '
                'checked against their PubMed records.'},
          {'h1': '2. Incretin physiology'},
          {'p': 'GIP is secreted by enteroendocrine K cells concentrated in the duodenum and proximal '
                'jejunum, whereas GLP-1 is secreted by L cells whose density increases toward the distal '
                'ileum and colon [[drucker18]]. Both hormones act through class B G protein-coupled '
                'receptors that couple predominantly to Gαs. Receptor activation raises intracellular cAMP '
                'and engages protein kinase A and Epac2; in β cells this amplifies insulin secretion only '
                'when glucose is already elevated, which is why incretin-based drugs carry little intrinsic '
                'risk of hypoglycaemia [[drucker18]].'},
          {'p': 'The two receptors differ in tissue distribution and function. GLP-1R is expressed in β '
                'cells, the stomach and intestine, the heart and vasculature, the kidney and discrete brain '
                'regions; its activation stimulates insulin, suppresses glucagon, slows gastric emptying and '
                'reduces appetite [[drucker18]]. GIPR is expressed in islet β and α cells, adipocytes and '
                'the brain. GIP stimulates insulin and, unlike GLP-1, promotes glucagon secretion at low '
                'glucose; in adipose tissue it enhances blood flow and triglyceride storage [[samms20]].'},
          {'p': 'In T2D the two incretins diverge. Under hyperglycaemic clamp conditions, GLP-1 retained '
                'most of its insulinotropic activity and also lowered glucagon, whereas the maximal insulin '
                'response to GIP was 54% lower than in healthy controls and GIP did not lower glucagon '
                '[[nauck93]]. This finding long discouraged GIP-based therapy and explains why the success '
                'of a GIP-dominant dual agonist was initially unexpected.'},
          {'h1': '3. Molecular pharmacology'},
          {'p': 'Both drugs achieve once-weekly dosing through the same two engineering steps: an '
                'α-aminoisobutyric acid (Aib) residue near the N-terminus blocks cleavage by dipeptidyl '
                'peptidase-4 (DPP-4), and a long fatty diacid side chain binds serum albumin, slowing renal '
                'clearance. They diverge in which receptors they engage and how they signal once bound '
                '(Table 1).'},
          {'h2': '3.1 Semaglutide'},
          {'p': 'Semaglutide is an analogue of human GLP-1 with two amino acid substitutions, Aib at '
                'position 8 and arginine at position 34, and a C18 fatty diacid attached to lysine 26 '
                'through a hydrophilic linker [[lau15]]. Its GLP-1R affinity (0.38 nM) is about three-fold '
                'lower than that of liraglutide, but its albumin affinity is higher; in mini-pigs the '
                'intravenous half-life was 46 hours [[lau15]], and in humans it is approximately one week '
                '[[marso16]]. Semaglutide is a full, selective GLP-1R agonist.'},
          {'h2': '3.2 Tirzepatide'},
          {'p': 'Tirzepatide is a 39-amino-acid peptide built on the native GIP sequence, carrying Aib '
                'residues and a C20 fatty diacid moiety that supports once-weekly dosing [[coskun18]]. It '
                'activates both GIPR and GLP-1R, with potency at GIPR similar to native GIP but lower '
                'potency at GLP-1R than native GLP-1 [[coskun18]]. Two features distinguish it from a simple '
                'combination of the two hormones [[willard20]]:'},
          {'bullets': ['**Imbalanced agonism.** Receptor-occupancy modelling at clinically effective doses '
                       'shows greater engagement of GIPR than of GLP-1R.',
                       '**Biased agonism at GLP-1R.** At GIPR, tirzepatide mimics native GIP. At GLP-1R, it '
                       'favours cAMP generation over β-arrestin recruitment and drives weaker receptor '
                       'internalisation than GLP-1. In primary islets, β-arrestin-1 limited the insulin '
                       'response to GLP-1 but not to tirzepatide, suggesting that the bias helps sustain '
                       'insulinotropic signalling.']},
          {'h2': '3.3 Structural basis'},
          {'p': 'Cryogenic electron microscopy and molecular dynamics simulations of tirzepatide bound to '
                'both receptors showed that it closely resembles native GIP in how it activates GIPR, but '
                'differs markedly from GLP-1 at GLP-1R, consistent with its reduced desensitisation of that '
                'receptor [[sun22]]. The fatty acid modification, together with the peptide sequence, '
                'contributes to this mode of action [[sun22]], which illustrates that lipidation is not '
                'merely a pharmacokinetic device.'},
          {'table': 't1'},
          {'h1': '4. Islet actions and insulin sensitivity'},
          {'p': 'Both drugs lower glucose primarily by restoring glucose-dependent insulin secretion and '
                'suppressing inappropriate glucagon release. Tirzepatide adds a measurable improvement in '
                'insulin sensitivity that is only partly explained by weight loss.'},
          {'p': 'In a 28-week mechanistic trial in T2D, tirzepatide 15 mg increased the clamp disposition '
                'index, an integrated measure of insulin secretion relative to sensitivity, significantly '
                'more than semaglutide 1 mg (estimated treatment difference 0.84; 95% CI 0.46–1.21) '
                '[[heise22]]. This reflected greater improvements in both total insulin secretion and '
                'clamp-derived insulin sensitivity (M value difference 1.52 mg/kg/min), and tirzepatide '
                'produced lower postprandial glucose, insulin and glucagon concentrations than semaglutide '
                'on meal testing [[heise22]].'},
          {'p': 'Preclinical work supports a direct GIPR contribution to insulin sensitisation. In obese '
                'mice lacking GLP-1R, and therefore without GLP-1R-driven weight loss, tirzepatide still '
                'improved insulin sensitivity by enhancing glucose disposal in white adipose tissue; a '
                'long-acting selective GIPR agonist reproduced this effect, which was accompanied by lower '
                'circulating branched-chain amino acids and ketoacids [[samms21]]. At the islet, the GIP '
                'component may also preserve counter-regulatory glucagon release as glucose falls '
                '[[samms20]].'},
          {'h1': '5. Central nervous system actions'},
          {'p': 'Weight loss with both drugs comes predominantly from eating less. In adults with obesity, '
                '20 weeks of semaglutide 2.4 mg reduced ad libitum energy intake by 35% compared with '
                'placebo, with less hunger, greater fullness, better control of eating and fewer and weaker '
                'food cravings [[friedrichsen21]]. In rodents, semaglutide reduced food intake and altered '
                'food preference without decreasing energy expenditure [[gabery20]].'},
          {'h2': '5.1 Access to the brain'},
          {'p': 'Albumin-bound semaglutide does not cross the intact blood–brain barrier. Fluorescently '
                'labelled semaglutide reached the brain through circumventricular organs and selected '
                'periventricular sites, including the brainstem, septal nucleus and hypothalamus '
                '[[gabery20]]. From these entry points it activated neurons in ten brain areas, including '
                'secondary regions without direct GLP-1R engagement such as the lateral parabrachial '
                'nucleus, a node implicated in meal termination [[gabery20]].'},
          {'h2': '5.2 Hypothalamus'},
          {'p': 'In the arcuate nucleus, labelled liraglutide was internalised by '
                'pro-opiomelanocortin/cocaine- and amphetamine-regulated transcript (POMC/CART) neurons. '
                'GLP-1 directly excited these anorexigenic neurons and indirectly inhibited orexigenic '
                'neuropeptide Y/agouti-related peptide (NPY/AgRP) neurons through GABA-dependent signalling '
                '[[secher14]]. Weight loss in rats persisted when GLP-1R was removed from vagal afferents, '
                'the area postrema or the paraventricular nucleus, pointing to arcuate POMC/CART neurons as '
                'a likely key site [[secher14]].'},
          {'h2': '5.3 Hindbrain and reward'},
          {'p': 'The area postrema (AP) and nucleus of the solitary tract (NTS) integrate vagal and '
                'circulating signals. Semaglutide increased expression of prolactin-releasing hormone and '
                'tyrosine hydroxylase in the AP [[gabery20]]. Hindbrain GLP-1R activation promotes satiation '
                'but also mediates nausea and conditioned aversion, linking efficacy and the main '
                'tolerability limit of the class. Beyond homeostatic feeding, GLP-1R signalling dampens '
                'reward. In a phase 2 trial in 48 adults with alcohol use disorder, low-dose semaglutide '
                'reduced laboratory alcohol self-administration, drinks per drinking day and craving '
                '[[hendershot25]].'},
          {'h2': '5.4 Central GIPR signalling'},
          {'p': 'Three lines of evidence suggest that the brain contributes to the added efficacy of '
                'tirzepatide. First, mice lacking GIPR in the central nervous system were resistant to the '
                'weight-lowering effect of acylated GIP, and the superiority of GIP/GLP-1 co-agonism over '
                'GLP-1 agonism alone was lost [[zhang21]]. Second, GIPR is expressed on distinct inhibitory '
                '(GABAergic) and excitatory neuron populations in the AP/NTS; in mice, rats and musk shrews, '
                'GIPR agonism blocked emesis and illness behaviours induced by GLP-1R agonists while '
                'preserving reduced food intake and weight loss [[borner21]]. Third, the picture is '
                'heterogeneous: hypothalamic GIPR was not required for the synergy of co-agonism on body '
                'weight, and chemogenetic activation of dorsal vagal complex GIPR neurons reduced intake but '
                'also induced conditioned taste avoidance [[adriaenssens23]]. Peripherally administered GIPR '
                'agonists reached only circumventricular organs [[adriaenssens23]].'},
          {'p': 'Together, these data support two complementary hypotheses: GIPR adds anorexigenic '
                'signalling, and GIPR improves tolerability, allowing greater effective GLP-1R engagement. '
                'Which predominates in humans is unresolved.'},
          {'h1': '6. Gastrointestinal actions'},
          {'p': 'GLP-1R activation slows gastric emptying through vagal pathways, flattening postprandial '
                'glucose excursions [[drucker18]]. With long-acting agonists this effect is largely lost '
                'over time. After 20 weeks of semaglutide 2.4 mg, gastric emptying assessed by paracetamol '
                'absorption showed no evidence of delay: early absorption measures were unchanged, and a '
                'small difference in total absorption over 5 hours was not significant after adjustment for '
                'body weight [[friedrichsen21]]. Gastric slowing is therefore unlikely to explain sustained '
                'weight loss, which was 9.9% in the same study [[friedrichsen21]].'},
          {'p': 'The same vagal and hindbrain pathways underlie the dominant adverse effects of the class: '
                'nausea, vomiting, diarrhoea and constipation, which cluster during dose escalation (Section '
                '10).'},
          {'figure': 'f1'},
          {'h1': '7. Adipose, liver, cardiovascular and renal effects'},
          {'h2': '7.1 Adipose tissue'},
          {'p': 'GLP-1R expression in adipocytes is sparse, whereas GIPR is abundant. GIP increases adipose '
                'blood flow and promotes lipoprotein lipase-mediated triglyceride uptake. The '
                'lipid-buffering hypothesis proposes that improved subcutaneous storage reduces spill-over '
                'of lipid into liver and muscle, improving insulin sensitivity [[samms20]]. The '
                'weight-independent insulin sensitisation seen in GLP-1R-null mice supports this model '
                '[[samms21]].'},
          {'h2': '7.2 Liver'},
          {'p': 'Hepatocytes express little or no GLP-1R, so liver benefits are thought to be largely '
                'indirect, reflecting weight loss, improved insulin sensitivity and reduced lipotoxicity. In '
                'the phase 3 ESSENCE trial in biopsy-proven metabolic dysfunction-associated steatohepatitis '
                '(MASH) with stage F2–F3 fibrosis, semaglutide 2.4 mg achieved resolution of steatohepatitis '
                'in 62.9% of patients versus 34.3% with placebo, and fibrosis improvement in 36.8% versus '
                '22.4%, at 72 weeks [[sanyal25]]. In the phase 2 SYNERGY-NASH trial, tirzepatide resolved '
                'MASH in 44–62% of participants across doses versus 10% with placebo at 52 weeks '
                '[[loomba24]].'},
          {'h2': '7.3 Heart and vasculature'},
          {'p': 'GLP-1R is expressed in the sinoatrial node, vascular smooth muscle and immune cells. '
                'Agonism lowers blood pressure, increases resting heart rate and reduces systemic '
                'inflammation [[drucker18]]. In SUSTAIN-6, semaglutide reduced major adverse cardiovascular '
                'events (MACE) in T2D (hazard ratio [HR] 0.74; 95% CI 0.58–0.95) [[marso16]]. In SELECT, '
                'which enrolled 17,604 people with established cardiovascular disease and overweight or '
                'obesity but without diabetes, semaglutide 2.4 mg reduced MACE by 20% (HR 0.80; 95% CI '
                '0.72–0.90) [[lincoff23]].'},
          {'p': 'Weight loss does not fully explain this benefit. In a prespecified SELECT analysis, MACE '
                'reduction was consistent across all categories of baseline weight and waist circumference; '
                'weight loss at week 20 showed no linear relationship with subsequent MACE risk; and an '
                'estimated 33% of the benefit was statistically mediated by reduction in waist circumference '
                '[[deanfield25]]. The remainder is attributed to direct vascular, anti-inflammatory and '
                'haemodynamic effects. In heart failure with preserved ejection fraction (HFpEF) and '
                'obesity, semaglutide improved symptoms and physical limitation and reduced C-reactive '
                'protein by 43.5% versus 7.3% with placebo [[kosiborod23]], while tirzepatide reduced the '
                'composite of cardiovascular death or worsening heart failure (HR 0.62; 95% CI 0.41–0.95) '
                '[[packer25]]. In SURPASS-CVOT, tirzepatide was non-inferior but not superior to dulaglutide '
                'for MACE (HR 0.92; 95.3% CI 0.83–1.01) [[nicholls25]].'},
          {'h2': '7.4 Kidney'},
          {'p': 'In FLOW, semaglutide 1 mg reduced major kidney disease events by 24% in people with T2D and '
                'chronic kidney disease (HR 0.76; 95% CI 0.66–0.88), slowed the annual decline in estimated '
                'glomerular filtration rate (eGFR) by 1.16 mL/min/1.73 m², and reduced cardiovascular death '
                '(HR 0.71) and death from any cause (HR 0.80) [[perkovic24]]. Proposed mechanisms include '
                'reduced glomerular hyperfiltration, natriuresis, lower blood pressure and anti-inflammatory '
                'effects.'},
          {'table': 't2'},
          {'h1': '8. The GIP paradox'},
          {'p': 'A central unresolved question is why both activating and blocking GIPR, when combined with '
                'GLP-1R agonism, enhance weight loss. Genome-wide association studies implicate GIPR in '
                'obesity, and GIPR knockout mice resist diet-induced obesity [[killion18]]. Antagonistic '
                'anti-GIPR antibodies protected obese mice and non-human primates against weight gain, and '
                'weight loss was enhanced when they were co-administered with GLP-1R agonists [[killion18]]. '
                'Maridebart cafraglutide (AMG 133), a GIPR-antagonist antibody conjugated to two GLP-1 '
                'analogue peptides, produced dose-dependent weight loss in a phase 1 trial that was '
                'maintained for up to 150 days after the last dose [[veniant24]].'},
          {'p': 'Several hypotheses attempt to reconcile these observations:'},
          {'bullets': ['**Receptor desensitisation.** Chronic GIPR agonism desensitised GIPR in primary '
                       'adipocytes in vitro and in adipose tissue in vivo, so a long-acting agonist '
                       'functioned like an antagonist in that tissue [[killion20]].',
                       '**Tissue-specific effects.** Agonism may dominate in the central nervous system, '
                       'adding anorexia and suppressing emesis [[zhang21,borner21]], while antagonist-like '
                       'effects dominate in adipose tissue.',
                       '**Different routes to the same end.** Agonism may act chiefly through tolerability '
                       'and adipose insulin sensitisation, and antagonism chiefly by amplifying '
                       'GLP-1R-driven anorexia.']},
          {'p': 'Direct mechanistic comparisons in humans are needed. Because tirzepatide is GIPR-favoured '
                '[[willard20]], its clinical profile cannot be read simply as GLP-1R agonism with a small '
                'GIP contribution.'},
          {'h1': '9. Clinical efficacy linked to mechanism'},
          {'p': 'In T2D, SURPASS-2 compared tirzepatide 5, 10 and 15 mg with semaglutide 1 mg over 40 weeks '
                'in 1,879 patients: HbA1c fell by 2.01, 2.24 and 2.30 percentage points with tirzepatide '
                'versus 1.86 points with semaglutide, and weight fell 1.9–5.5 kg more with tirzepatide '
                '[[frias21]]. In obesity without diabetes, semaglutide 2.4 mg reduced weight by 14.9% versus '
                '2.4% with placebo at 68 weeks (STEP 1) [[wilding21]], and tirzepatide 15 mg by 20.9% versus '
                '3.1% at 72 weeks (SURMOUNT-1) [[jastreboff22]]. In the head-to-head SURMOUNT-5 trial, '
                'tirzepatide reduced weight by 20.2% versus 13.7% with semaglutide at 72 weeks, with greater '
                'waist reduction (18.4 vs 13.0 cm) [[aronne25]] (Figure 2).'},
          {'figure': 'f2'},
          {'p': "Three mechanistic explanations for tirzepatide's advantage are compatible with these data: "
                'additional central anorexigenic signalling via GIPR [[zhang21]]; better tolerability '
                'through hindbrain GIPR, allowing greater effective GLP-1R engagement [[borner21]]; and '
                'GIPR-mediated adipose insulin sensitisation [[heise22,samms21]]. Cross-trial comparisons '
                'are limited by differences in populations, dose escalation and estimands, and SURMOUNT-5 '
                'was open-label.'},
          {'h2': '9.1 Durability and the chronic nature of the mechanism'},
          {'p': 'Because these drugs act on homeostatic and reward circuits rather than permanently altering '
                'them, their effects depend on continued exposure. One year after stopping semaglutide and '
                'lifestyle intervention in the STEP 1 extension, participants regained two-thirds of their '
                'prior weight loss, and cardiometabolic improvements reverted toward baseline [[wilding22]]. '
                'In SURMOUNT-4, after a 36-week tirzepatide lead-in producing 20.9% weight loss, '
                'participants switched to placebo regained 14.0% of body weight over 52 weeks, whereas those '
                'who continued tirzepatide lost a further 5.5% [[aronne24]].'},
          {'h1': '10. Safety and mechanism-linked adverse effects'},
          {'p': 'Most adverse effects of both drugs are predictable extensions of incretin receptor '
                'pharmacology (Table 3). Gastrointestinal events dominate: in SURPASS-2, nausea occurred in '
                '17–22% of tirzepatide recipients and 18% of semaglutide recipients [[frias21]]. '
                'Discontinuation because of adverse events in SURMOUNT-1 was 4.3–7.1% with tirzepatide '
                'versus 2.6% with placebo [[jastreboff22]], and in STEP 1, 4.5% of semaglutide recipients '
                'stopped because of gastrointestinal events versus 0.8% on placebo [[wilding21]].'},
          {'p': 'Lean mass is a growing concern. A network meta-analysis of 22 randomised trials found that '
                'lean mass accounted for about 25% of weight lost with GLP-1R-based agents; tirzepatide 15 '
                'mg and semaglutide 2.4 mg were the most effective at reducing fat but among the least '
                'lean-sparing [[karakasis24]]. This has prompted trials combining incretin drugs with '
                'inhibitors of myostatin and activin signalling, such as bimagrumab [[stefanakis24]].'},
          {'p': 'In rodents, GLP-1R agonists activate thyroid C cells, causing calcitonin release and C-cell '
                'hyperplasia. Primate C cells express little GLP-1R and did not respond, and 20 months of '
                'high-dose liraglutide did not cause C-cell hyperplasia in monkeys [[bjerre10]]. Human '
                'relevance is therefore uncertain, but both drugs are contraindicated in people with a '
                'personal or family history of medullary thyroid carcinoma or multiple endocrine neoplasia '
                'type 2. Rapid glycaemic improvement was associated with more diabetic retinopathy '
                'complications in SUSTAIN-6 (HR 1.76; 95% CI 1.11–2.78) [[marso16]]. In a single-centre '
                'retrospective cohort, semaglutide was associated with a higher 36-month cumulative '
                'incidence of non-arteritic anterior ischaemic optic neuropathy (NAION) in patients with T2D '
                '(8.9% vs 1.8%) [[hathaway24]]; causality and mechanism remain unresolved.'},
          {'table': 't3'},
          {'h1': '11. Knowledge gaps and future directions'},
          {'p': 'The field is moving from one or two receptors to several, and from injected peptides to '
                'oral small molecules. Retatrutide, which adds glucagon receptor agonism to GIPR and GLP-1R '
                'agonism, reduced weight by 24.2% at 48 weeks with the 12 mg dose in phase 2 '
                '[[jastreboff23]]. CagriSema, which combines semaglutide with the long-acting amylin '
                'analogue cagrilintide, reduced weight by 20.4% versus 3.0% with placebo at 68 weeks in '
                'REDEFINE 1 [[garvey25]]. Orforglipron, a once-daily oral non-peptide GLP-1R agonist, '
                'reduced weight by 11.2% with the 36 mg dose versus 2.1% with placebo at 72 weeks '
                '[[wharton25]]. Further clinical trials of GIPR antagonism with maridebart cafraglutide will '
                'directly test the GIP paradox [[veniant24]].'},
          {'p': 'Key mechanistic questions remain open:'},
          {'bullets': ['Which tissues mediate the GIPR contribution in humans, and does sustained agonism '
                       'act as functional antagonism?',
                       'How much of the cardiovascular and renal benefit is independent of weight loss, and '
                       'is it shared by tirzepatide?',
                       'Can lean mass be preserved without sacrificing fat loss?',
                       'How can benefit be maintained after treatment is reduced or stopped?',
                       'What is the long-term safety of chronic multi-receptor agonism, including ocular and '
                       'pancreatobiliary effects?']},
          {'h1': '12. Conclusions'},
          {'p': 'Semaglutide acts through potent, sustained and selective GLP-1R activation across the '
                'islet, gut, brain, heart and kidney. Tirzepatide engages the same pathways through a '
                'GIPR-favoured, cAMP-biased dual agonist profile that adds central GIPR signalling, '
                'potential tolerability benefits and adipose-mediated insulin sensitisation. These '
                "differences are consistent with tirzepatide's greater glycaemic and weight effects, whereas "
                'semaglutide currently has the broader base of proven cardiovascular and kidney outcomes. '
                'For both drugs, the dominant mechanism of weight loss is central suppression of food '
                'intake, which is why benefits persist only with continued treatment. Resolving the GIP '
                'paradox and the weight-independent mechanisms of organ protection will shape the next '
                'generation of incretin-based medicines.'}],
 'TABLES': {'t1': {'caption': 'Table 1. Molecular and pharmacological comparison of semaglutide and '
                              'tirzepatide',
                   'widths': [2500, 3430, 3430],
                   'head': ['Feature', 'Semaglutide', 'Tirzepatide'],
                   'rows': [['Brand names',
                             'Ozempic (T2D), Wegovy (obesity), Rybelsus (oral, T2D)',
                             'Mounjaro (T2D), Zepbound (obesity)'],
                            ['Receptor targets', 'GLP-1R', 'GIPR and GLP-1R'],
                            ['Peptide backbone',
                             'Human GLP-1(7–37) analogue',
                             'Native GIP sequence, 39 amino acids'],
                            ['DPP-4 protection', 'Aib at position 8', 'Aib residues'],
                            ['Albumin binder', 'C18 fatty diacid on Lys26', 'C20 fatty diacid'],
                            ['Half-life in humans', 'About 1 week', 'About 5 days*'],
                            ['Receptor balance', 'Selective GLP-1R', 'Imbalanced: greater GIPR engagement'],
                            ['GLP-1R signalling',
                             'Full agonist',
                             'cAMP-biased; less β-arrestin recruitment and internalisation'],
                            ['Weekly maintenance doses',
                             '0.5–2 mg (T2D); 2.4 mg (obesity)',
                             '5, 10 or 15 mg']],
                   'note': 'Sources: [[lau15,marso16,coskun18,willard20,sun22]]. *From the US prescribing '
                           'information. Aib, α-aminoisobutyric acid; DPP-4, dipeptidyl peptidase-4.'},
            't2': {'caption': 'Table 2. Cardiovascular, kidney and liver outcome trials',
                   'widths': [1500, 2900, 1700, 3260],
                   'head': ['Trial', 'Population', 'Comparison', 'Primary result'],
                   'rows': [['SUSTAIN-6 [[marso16]]',
                             'T2D at high CV risk (n = 3,297)',
                             'Semaglutide vs placebo',
                             'MACE HR 0.74 (0.58–0.95)'],
                            ['SELECT [[lincoff23]]',
                             'CVD with overweight or obesity, no diabetes (n = 17,604)',
                             'Semaglutide 2.4 mg vs placebo',
                             'MACE HR 0.80 (0.72–0.90)'],
                            ['FLOW [[perkovic24]]',
                             'T2D with CKD (n = 3,533)',
                             'Semaglutide 1 mg vs placebo',
                             'Major kidney disease events HR 0.76 (0.66–0.88)'],
                            ['STEP-HFpEF [[kosiborod23]]',
                             'HFpEF with obesity (n = 529)',
                             'Semaglutide 2.4 mg vs placebo',
                             'KCCQ-CSS +7.8 points; weight −10.7 percentage points'],
                            ['SUMMIT [[packer25]]',
                             'HFpEF with obesity (n = 731)',
                             'Tirzepatide vs placebo',
                             'CV death or worsening HF HR 0.62 (0.41–0.95)'],
                            ['SURPASS-CVOT [[nicholls25]]',
                             'T2D with ASCVD (n = 13,165 analysed)',
                             'Tirzepatide vs dulaglutide',
                             'MACE HR 0.92 (0.83–1.01); non-inferior'],
                            ['ESSENCE [[sanyal25]]',
                             'MASH, F2–F3 fibrosis (n = 800, interim)',
                             'Semaglutide 2.4 mg vs placebo',
                             'MASH resolution 62.9% vs 34.3%'],
                            ['SYNERGY-NASH [[loomba24]]',
                             'MASH, F2–F3 fibrosis (n = 190); phase 2',
                             'Tirzepatide vs placebo',
                             'MASH resolution 44–62% vs 10%']],
                   'note': 'Hazard ratios with 95% confidence intervals (95.3% for SURPASS-CVOT). ASCVD, '
                           'atherosclerotic cardiovascular disease; CKD, chronic kidney disease; KCCQ-CSS, '
                           'Kansas City Cardiomyopathy Questionnaire clinical summary score.'},
            't3': {'caption': 'Table 3. Mechanism-linked adverse effects',
                   'widths': [2300, 3800, 3260],
                   'head': ['Adverse effect', 'Proposed mechanism', 'Evidence'],
                   'rows': [['Nausea, vomiting, diarrhoea',
                             'Hindbrain (AP/NTS) and vagal GLP-1R activation',
                             'Most common events; peak during dose escalation [[frias21,wilding21]]'],
                            ['Lean mass loss',
                             'Negative energy balance; magnitude of weight loss',
                             'About 25% of weight lost [[karakasis24]]'],
                            ['Diabetic retinopathy worsening',
                             'Rapid fall in HbA1c',
                             'SUSTAIN-6 HR 1.76 (1.11–2.78) [[marso16]]'],
                            ['Thyroid C-cell hyperplasia',
                             'Dense C-cell GLP-1R in rodents, sparse in primates',
                             'Rodent-specific to date [[bjerre10]]'],
                            ['Increased heart rate',
                             'Sinoatrial node GLP-1R',
                             'Consistent small increase [[drucker18]]'],
                            ['NAION', 'Unknown', 'Observational association only [[hathaway24]]']],
                   'note': 'AP, area postrema; NAION, non-arteritic anterior ischaemic optic neuropathy; '
                           'NTS, nucleus of the solitary tract.'}},
 'FIGURES': {'f1': {'file': 'fig1.png',
                    'w': 620,
                    'h': 496,
                    'caption': '**Figure 1. Integrated sites of action of semaglutide and tirzepatide.** '
                               'Both drugs act through GLP-1R in the brain, islet, gut, heart and kidney '
                               '(blue). Tirzepatide additionally engages GIPR (orange), adding central '
                               'anorexigenic and anti-emetic signalling, insulinotropic effects and adipose '
                               'insulin sensitisation. Liver benefits are largely indirect. Based on '
                               '[[drucker18,samms20,willard20,gabery20,borner21,samms21]]; outcome trials as '
                               'cited in Table 2.'},
             'f2': {'file': 'fig2.png',
                    'w': 620,
                    'h': 296,
                    'caption': '**Figure 2. Mean percentage change in body weight in pivotal trials.** Each '
                               'row joins the active treatment to its comparator. SURMOUNT-5 is the only '
                               'head-to-head comparison of tirzepatide and semaglutide; other comparisons '
                               'across trials are indirect and should be interpreted cautiously. Data from '
                               '[[wilding21,jastreboff22,aronne25,wharton25,garvey25,jastreboff23]]. MTD, '
                               'maximum tolerated dose.'}}}

if __name__ == "__main__":
    main()
